#!/usr/bin/env python3
"""Generate an email-style static-depth ladder report from run bundles."""

from __future__ import annotations

import argparse
import csv
import html
import json
import math
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
from matplotlib import pyplot as plt  # noqa: E402

from docx import Document  # type: ignore
from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore
from docx.shared import Inches, Pt  # type: ignore


@dataclass
class RunSpec:
    label: str
    run_dir: Path
    section_title: str
    note: str
    figure_name: str


@dataclass
class RunMetrics:
    spec: RunSpec
    frame_rows_total: int
    track_rows_total: int
    frame_window_rows: int
    track_window_rows_matched: int
    truth_range_mean_m: float | None
    truth_depth_mean_m: float | None
    detection_fraction_window: float
    world_track_fraction_window: float
    controller_valid_fraction_window: float
    estimated_distance_mean_m: float | None
    estimated_depth_mean_m: float | None
    error_p90_m_window: float | None
    error_mean_m_window: float | None
    ready_for_follow_overall: bool | None
    representative_image_src: Path | None
    representative_image_dst: Path | None = None


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate an email-style static-depth ladder report."
    )
    parser.add_argument("--title", required=True)
    parser.add_argument("--topnote", required=True)
    parser.add_argument("--summary-heading", required=True)
    parser.add_argument("--coverage-caption", required=True)
    parser.add_argument("--error-caption", required=True)
    parser.add_argument("--writeup-path", required=True)
    parser.add_argument("--html-path", required=True)
    parser.add_argument("--docx-path", required=True)
    parser.add_argument("--odt-path", default="")
    parser.add_argument("--assets-dir", required=True)
    parser.add_argument("--email-frames-dir", required=True)
    parser.add_argument("--window-size-frames", type=int, default=180)
    parser.add_argument(
        "--run",
        action="append",
        default=[],
        help=(
            "Run spec encoded as "
            "'label|run_dir|section_title|note|figure_name'. Repeat once per run."
        ),
    )
    return parser.parse_args()


def parse_run_spec(value: str) -> RunSpec:
    parts = [part.strip() for part in value.split("|", 4)]
    if len(parts) != 5:
        raise ValueError(
            "Run spec must have exactly 5 pipe-delimited fields: "
            "label|run_dir|section_title|note|figure_name"
        )
    label, run_dir, section_title, note, figure_name = parts
    return RunSpec(
        label=label,
        run_dir=Path(run_dir),
        section_title=section_title,
        note=note,
        figure_name=figure_name,
    )


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def optional_float(value: str | None) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if not math.isfinite(number):
        return None
    return number


def optional_bool(value: str | None) -> bool | None:
    if value is None:
        return None
    text = str(value).strip().lower()
    if not text:
        return None
    if text in {"1", "true", "yes"}:
        return True
    if text in {"0", "false", "no"}:
        return False
    return None


def safe_mean(values: list[float]) -> float | None:
    if not values:
        return None
    return sum(values) / len(values)


def percentile(values: list[float], q: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    idx = (len(ordered) - 1) * q
    lo = int(math.floor(idx))
    hi = int(math.ceil(idx))
    if lo == hi:
        return ordered[lo]
    return ordered[lo] + (ordered[hi] - ordered[lo]) * (idx - lo)


def extract_series(rows: list[dict[str, str]], key: str) -> list[float]:
    values: list[float] = []
    for row in rows:
        number = optional_float(row.get(key))
        if number is not None:
            values.append(number)
    return values


def pick_representative_image(run_dir: Path) -> Path | None:
    images = sorted(run_dir.glob("world_track_compare_*.jpg"))
    if not images:
        return None
    return images[len(images) // 2]


def load_ready_for_follow(run_dir: Path) -> bool | None:
    summary_path = run_dir / "report" / "summary.json"
    if not summary_path.exists():
        return None
    data = json.loads(summary_path.read_text(encoding="utf-8"))
    follow_readiness = data.get("follow_readiness", {})
    if not isinstance(follow_readiness, dict):
        return None
    ready = follow_readiness.get("ready_for_follow")
    if isinstance(ready, bool):
        return ready
    return None


def summarize_run(spec: RunSpec, window_size_frames: int) -> RunMetrics:
    frames = read_csv_rows(spec.run_dir / "world_track_compare_frames.csv")
    tracks = read_csv_rows(spec.run_dir / "world_track_compare_tracks.csv")
    frame_window = (
        frames[-window_size_frames:] if len(frames) > window_size_frames else frames
    )
    start_time = optional_float(frame_window[0].get("timestamp_s")) if frame_window else None
    end_time = optional_float(frame_window[-1].get("timestamp_s")) if frame_window else None

    track_window: list[dict[str, str]] = []
    if start_time is not None and end_time is not None:
        for row in tracks:
            timestamp = optional_float(row.get("timestamp_s"))
            if timestamp is None:
                continue
            if start_time <= timestamp <= end_time:
                track_window.append(row)

    frame_window_len = len(frame_window)
    detection_fraction = (
        sum(
            1
            for row in frame_window
            if (optional_float(row.get("detections_count")) or 0.0) > 0.0
        )
        / frame_window_len
        if frame_window_len
        else 0.0
    )
    world_track_fraction = (
        sum(
            1
            for row in frame_window
            if (optional_float(row.get("world_tracks_count")) or 0.0) > 0.0
        )
        / frame_window_len
        if frame_window_len
        else 0.0
    )
    controller_valid_fraction = (
        sum(1 for row in frame_window if optional_bool(row.get("has_controller_valid_track")))
        / frame_window_len
        if frame_window_len
        else 0.0
    )

    error_values = extract_series(track_window, "error_norm_m")

    return RunMetrics(
        spec=spec,
        frame_rows_total=len(frames),
        track_rows_total=len(tracks),
        frame_window_rows=len(frame_window),
        track_window_rows_matched=len(track_window),
        truth_range_mean_m=safe_mean(extract_series(frame_window, "truth_range_m")),
        truth_depth_mean_m=safe_mean(extract_series(frame_window, "truth_depth_z_m")),
        detection_fraction_window=detection_fraction,
        world_track_fraction_window=world_track_fraction,
        controller_valid_fraction_window=controller_valid_fraction,
        estimated_distance_mean_m=safe_mean(extract_series(track_window, "distance_m")),
        estimated_depth_mean_m=safe_mean(extract_series(track_window, "depth_z_m")),
        error_p90_m_window=percentile(error_values, 0.9),
        error_mean_m_window=safe_mean(error_values),
        ready_for_follow_overall=load_ready_for_follow(spec.run_dir),
        representative_image_src=pick_representative_image(spec.run_dir),
    )


def ensure_clean_dir(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True, exist_ok=True)


def format_float(value: float | None, digits: int = 3) -> str:
    if value is None:
        return "n/a"
    return f"{value:.{digits}f}"


def format_percent(value: float | None) -> str:
    if value is None:
        return "n/a"
    return f"{value * 100.0:.1f}%"


def write_summary_csv(metrics: list[RunMetrics], summary_csv_path: Path) -> None:
    summary_csv_path.parent.mkdir(parents=True, exist_ok=True)
    with summary_csv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "label",
                "run_id",
                "note",
                "frame_rows_total",
                "track_rows_total",
                "frame_window_rows",
                "track_window_rows_matched",
                "truth_range_mean_m",
                "truth_depth_mean_m",
                "detection_fraction_window",
                "world_track_fraction_window",
                "controller_valid_fraction_window",
                "estimated_distance_mean_m",
                "estimated_depth_mean_m",
                "error_p90_m_window",
                "error_mean_m_window",
                "ready_for_follow_overall",
            ]
        )
        for item in metrics:
            writer.writerow(
                [
                    item.spec.label,
                    item.spec.run_dir.name,
                    item.spec.note,
                    item.frame_rows_total,
                    item.track_rows_total,
                    item.frame_window_rows,
                    item.track_window_rows_matched,
                    item.truth_range_mean_m,
                    item.truth_depth_mean_m,
                    item.detection_fraction_window,
                    item.world_track_fraction_window,
                    item.controller_valid_fraction_window,
                    item.estimated_distance_mean_m,
                    item.estimated_depth_mean_m,
                    item.error_p90_m_window,
                    item.error_mean_m_window,
                    item.ready_for_follow_overall,
                ]
            )


def make_line_plot(
    metrics: list[RunMetrics],
    *,
    y_getter,
    ylabel: str,
    title: str,
    output_path: Path,
    color: str,
) -> None:
    xs: list[float] = []
    ys: list[float] = []
    labels: list[str] = []
    for item in metrics:
        x = item.truth_range_mean_m
        y = y_getter(item)
        if x is None or y is None:
            continue
        xs.append(x)
        ys.append(y)
        labels.append(item.spec.label)

    plt.figure(figsize=(8.0, 4.6))
    if xs:
        plt.plot(xs, ys, marker="o", linewidth=2.0, color=color)
        for x, y, label in zip(xs, ys, labels, strict=True):
            plt.annotate(
                label,
                (x, y),
                xytext=(0, 8),
                textcoords="offset points",
                ha="center",
                fontsize=9,
            )
    plt.xlabel("Ground-truth range in final steady window (m)")
    plt.ylabel(ylabel)
    plt.title(title)
    plt.grid(True, alpha=0.3)
    if "fraction" in ylabel.lower():
        plt.ylim(-0.05, 1.05)
    plt.tight_layout()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    plt.savefig(output_path, dpi=180)
    plt.close()


def copy_report_assets(
    metrics: list[RunMetrics],
    *,
    email_frames_dir: Path,
    coverage_plot_src: Path,
    error_plot_src: Path,
) -> tuple[Path, Path]:
    ensure_clean_dir(email_frames_dir)
    coverage_dst = email_frames_dir / "07_range_vs_tracking_fraction.png"
    error_dst = email_frames_dir / "08_range_vs_error_p90.png"
    shutil.copy2(coverage_plot_src, coverage_dst)
    shutil.copy2(error_plot_src, error_dst)

    for item in metrics:
        if item.representative_image_src is None:
            continue
        dst_path = email_frames_dir / item.spec.figure_name
        shutil.copy2(item.representative_image_src, dst_path)
        item.representative_image_dst = dst_path

    return coverage_dst, error_dst


def build_run_caption(item: RunMetrics) -> str:
    prefix = item.spec.note.strip()
    if prefix and not prefix.endswith((".", "!", "?")):
        prefix = prefix + "."

    common = (
        f"{prefix} Ground-truth range: {format_float(item.truth_range_mean_m)} m. "
        f"Ground-truth forward depth: {format_float(item.truth_depth_mean_m)} m. "
        f"Detector hit rate in the final steady window: "
        f"{format_percent(item.detection_fraction_window)}. "
        f"World-track hit rate: {format_percent(item.world_track_fraction_window)}. "
        f"Follow-usable hit rate: {format_percent(item.controller_valid_fraction_window)}."
    )

    if item.error_p90_m_window is None:
        return (
            f"{common} There was no estimated depth and no position error in this "
            "window because the detector never produced a valid world track."
        )

    return (
        f"{common} Mean estimated depth from the tracker: "
        f"{format_float(item.estimated_depth_mean_m)} m. "
        f"P90 world-position error: {format_float(item.error_p90_m_window)} m."
    )


def write_html_report(
    *,
    title: str,
    topnote: str,
    coverage_title: str,
    coverage_image: Path,
    coverage_caption: str,
    error_title: str,
    error_image: Path,
    error_caption: str,
    metrics: list[RunMetrics],
    html_path: Path,
) -> None:
    parts = [
        "<!DOCTYPE html>",
        "<html><head><meta charset=\"utf-8\">",
        f"<title>{html.escape(title)}</title>",
        "<style>",
        "body { font-family: Arial, sans-serif; margin: 0.75in; color: #111; line-height: 1.35; }",
        "h1 { font-size: 22pt; margin: 0 0 8pt 0; }",
        "p.topnote { font-size: 11pt; margin: 0 0 16pt 0; }",
        "h2 { font-size: 15pt; margin: 24pt 0 8pt 0; page-break-after: avoid; }",
        "img.figure { width: 6.5in; max-width: 100%; display: block; margin: 8pt 0 8pt 0; border: 1px solid #ccc; }",
        "p.caption { font-size: 11pt; margin: 0 0 18pt 0; }",
        "</style></head><body>",
        f"<h1>{html.escape(title)}</h1>",
        f"<p class=\"topnote\">{html.escape(topnote)}</p>",
        f"<h2>{html.escape(coverage_title)}</h2>",
        (
            f"<img class=\"figure\" src=\"file://{coverage_image}\" "
            f"alt=\"{html.escape(coverage_title)}\">"
        ),
        (
            f"<p class=\"caption\">{html.escape(coverage_caption)}"
            f"<br><span style=\"font-size:9pt;color:#555;\">"
            f"Source image: {html.escape(str(coverage_image))}</span></p>"
        ),
        f"<h2>{html.escape(error_title)}</h2>",
        (
            f"<img class=\"figure\" src=\"file://{error_image}\" "
            f"alt=\"{html.escape(error_title)}\">"
        ),
        (
            f"<p class=\"caption\">{html.escape(error_caption)}"
            f"<br><span style=\"font-size:9pt;color:#555;\">"
            f"Source image: {html.escape(str(error_image))}</span></p>"
        ),
    ]

    for item in metrics:
        if item.representative_image_dst is None:
            continue
        caption = build_run_caption(item)
        parts.extend(
            [
                f"<h2>{html.escape(item.spec.section_title)}</h2>",
                (
                    f"<img class=\"figure\" src=\"file://{item.representative_image_dst}\" "
                    f"alt=\"{html.escape(item.spec.section_title)}\">"
                ),
                (
                    f"<p class=\"caption\">{html.escape(caption)}"
                    f"<br><span style=\"font-size:9pt;color:#555;\">"
                    f"Source image: {html.escape(str(item.representative_image_dst))}"
                    f"</span></p>"
                ),
            ]
        )

    parts.append("</body></html>")
    html_path.write_text("\n".join(parts), encoding="utf-8")


def add_picture_paragraph(document: Document, image_path: Path) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.5))


def add_caption_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.add_run(text)


def write_docx_report(
    *,
    title: str,
    topnote: str,
    coverage_title: str,
    coverage_image: Path,
    coverage_caption: str,
    error_title: str,
    error_image: Path,
    error_caption: str,
    metrics: list[RunMetrics],
    docx_path: Path,
    email_frames_dir: Path,
) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(20)

    topnote_para = document.add_paragraph()
    topnote_prefix = topnote_para.add_run("Image-first summary for email. ")
    topnote_prefix.bold = True
    topnote_para.add_run(topnote)

    overview_heading = document.add_paragraph()
    overview_heading.add_run(coverage_title).bold = True
    overview_heading.runs[0].font.size = Pt(14)
    add_picture_paragraph(document, coverage_image)
    add_caption_paragraph(document, coverage_caption)

    error_heading = document.add_paragraph()
    error_heading.add_run(error_title).bold = True
    error_heading.runs[0].font.size = Pt(14)
    add_picture_paragraph(document, error_image)
    add_caption_paragraph(document, error_caption)

    for item in metrics:
        if item.representative_image_dst is None:
            continue
        document.add_page_break()
        heading = document.add_paragraph()
        heading_run = heading.add_run(item.spec.section_title)
        heading_run.bold = True
        heading_run.font.size = Pt(14)
        add_picture_paragraph(document, item.representative_image_dst)
        add_caption_paragraph(document, build_run_caption(item))

    source_para = document.add_paragraph()
    source_prefix = source_para.add_run("Embedded source images are also in ")
    source_prefix.bold = True
    source_para.add_run(str(email_frames_dir))
    document.save(str(docx_path))


def convert_docx_to_odt(docx_path: Path, odt_path: Path) -> None:
    odt_path.parent.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to",
            "odt",
            "--outdir",
            str(odt_path.parent),
            str(docx_path),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    converted = odt_path.parent / (docx_path.stem + ".odt")
    if converted != odt_path:
        shutil.move(converted, odt_path)
    if result.stdout.strip():
        print(result.stdout.strip())
    if result.stderr.strip():
        print(result.stderr.strip())


def write_markdown_writeup(
    *,
    title: str,
    summary_heading: str,
    topnote: str,
    metrics: list[RunMetrics],
    writeup_path: Path,
    summary_csv_path: Path,
    coverage_plot_path: Path,
    error_plot_path: Path,
) -> None:
    lines = [
        f"# {title}",
        "",
        "## Purpose",
        "",
        topnote,
        "",
        f"## {summary_heading}",
        "",
        "| Nominal hold | Run ID | Ground-truth range mean (m) | Ground-truth depth mean (m) | Detection fraction | World-track fraction | Controller-valid fraction | Estimated depth mean (m) | P90 error (m) | Notes |",
        "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- |",
    ]

    for item in metrics:
        lines.append(
            "| "
            + " | ".join(
                [
                    item.spec.label,
                    f"`{item.spec.run_dir.name}`",
                    format_float(item.truth_range_mean_m),
                    format_float(item.truth_depth_mean_m),
                    f"{item.detection_fraction_window:.3f}",
                    f"{item.world_track_fraction_window:.3f}",
                    f"{item.controller_valid_fraction_window:.3f}",
                    format_float(item.estimated_depth_mean_m),
                    format_float(item.error_p90_m_window),
                    item.spec.note,
                ]
            )
            + " |"
        )

    lines.extend(
        [
            "",
            "Source CSV for this table:",
            "",
            f"- `{summary_csv_path}`",
            "",
            "Overview plots:",
            "",
            f"- Tracking coverage vs range: `{coverage_plot_path}`",
            f"- Error vs range: `{error_plot_path}`",
            "",
            "Representative runs:",
            "",
        ]
    )

    for item in metrics:
        lines.append(
            f"- {item.spec.label}: `{item.spec.run_dir / 'report' / 'summary.md'}`"
        )

    writeup_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    args = parse_args()
    run_specs = [parse_run_spec(value) for value in args.run]
    if not run_specs:
        raise SystemExit("At least one --run spec is required.")

    metrics = [summarize_run(spec, args.window_size_frames) for spec in run_specs]

    assets_dir = Path(args.assets_dir)
    email_frames_dir = Path(args.email_frames_dir)
    assets_dir.mkdir(parents=True, exist_ok=True)

    summary_csv_path = assets_dir / "static_depth_ladder_window_summary.csv"
    coverage_plot_path = assets_dir / "range_vs_tracking_fraction.png"
    error_plot_path = assets_dir / "range_vs_error_p90.png"

    write_summary_csv(metrics, summary_csv_path)
    make_line_plot(
        metrics,
        y_getter=lambda item: item.world_track_fraction_window,
        ylabel="World-track fraction in final steady window",
        title="Tracking Coverage vs Ground-Truth Range",
        output_path=coverage_plot_path,
        color="#0b6e4f",
    )
    make_line_plot(
        metrics,
        y_getter=lambda item: item.error_p90_m_window,
        ylabel="P90 world-position error in final steady window (m)",
        title="Error vs Ground-Truth Range",
        output_path=error_plot_path,
        color="#aa3a1e",
    )

    coverage_image, error_image = copy_report_assets(
        metrics,
        email_frames_dir=email_frames_dir,
        coverage_plot_src=coverage_plot_path,
        error_plot_src=error_plot_path,
    )

    html_path = Path(args.html_path)
    docx_path = Path(args.docx_path)
    writeup_path = Path(args.writeup_path)

    write_html_report(
        title=args.title,
        topnote=args.topnote,
        coverage_title="Overview Plot: Tracking Coverage vs Ground-Truth Range",
        coverage_image=coverage_image,
        coverage_caption=args.coverage_caption,
        error_title="Overview Plot: Error vs Ground-Truth Range",
        error_image=error_image,
        error_caption=args.error_caption,
        metrics=metrics,
        html_path=html_path,
    )
    write_docx_report(
        title=args.title,
        topnote=args.topnote,
        coverage_title="Overview Plot: Tracking Coverage vs Ground-Truth Range",
        coverage_image=coverage_image,
        coverage_caption=args.coverage_caption,
        error_title="Overview Plot: Error vs Ground-Truth Range",
        error_image=error_image,
        error_caption=args.error_caption,
        metrics=metrics,
        docx_path=docx_path,
        email_frames_dir=email_frames_dir,
    )
    write_markdown_writeup(
        title=args.title,
        summary_heading=args.summary_heading,
        topnote=args.topnote,
        metrics=metrics,
        writeup_path=writeup_path,
        summary_csv_path=summary_csv_path,
        coverage_plot_path=coverage_plot_path,
        error_plot_path=error_plot_path,
    )

    if args.odt_path:
        convert_docx_to_odt(docx_path, Path(args.odt_path))

    print(f"Wrote markdown writeup: {writeup_path}")
    print(f"Wrote HTML report: {html_path}")
    print(f"Wrote DOCX report: {docx_path}")
    print(f"Wrote assets directory: {assets_dir}")
    print(f"Wrote email frames directory: {email_frames_dir}")


if __name__ == "__main__":
    main()
